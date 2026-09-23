# import os
# import re
# from docx import Document
# from langchain.memory import ConversationBufferMemory
# from langchain.chains import ConversationChain
# from langchain.chat_models import ChatOpenAI

# # === 🔑 Set your OpenAI API key here ===
# os.environ["OPENAI_API_KEY"] = ""  # Replace with your actual key

# # === 🧠 Track conversation state ===
# user_state = {
#     "Intent": None,
#     "User_type": None,
#     "property_type": None,
#     "sub_type": None,
#     "location": None,
#     "budget": None,
#     "bhk": None,
#     "sqft": None,
#     "furnishing": None,
# }

# # === 📄 Load prompts from Word document ===
# def load_prompt_response_table(docx_path):
#     doc = Document(docx_path)
#     prompt_map = {}
#     for table in doc.tables:
#         for row in table.rows[1:]:
#             prompt = row.cells[0].text.strip().lower()
#             response = "\n".join(
#                 para.text.strip() for para in row.cells[1].paragraphs if para.text.strip()
#             )
#             prompt_map[prompt] = response
#     return prompt_map

# # === 🤖 Create OpenAI LLM chatbot with memory ===
# def create_llm_chatbot():
#     llm = ChatOpenAI(model="gpt-3.5-turbo", temperature=0.5)
#     memory = ConversationBufferMemory()
#     return ConversationChain(llm=llm, memory=memory, verbose=False)

# # === 🧠 Use LLM for fuzzy prompt matching ===
# def classify_to_prompt(user_input, prompt_map, llm):
#     prompt_text = (
#         "You are an assistant that classifies user input into one of the following known prompts.\n"
#         "Known prompts:\n" +
#         "\n".join(f"- {p}" for p in prompt_map.keys()) +
#         f"\n\nUser input: {user_input}\n\n"
#         "Return ONLY the exact matching prompt from the list, or 'unknown' if it doesn't match."
#     )
#     try:
#         result = llm.invoke(prompt_text)
#         matched_prompt = result.content.strip().lower()
#         return prompt_map.get(matched_prompt, None)
#     except Exception as e:
#         print("❌ LLM Classification Error:", str(e))
#         return None

# # === 🏙️ Detect location input ===
# def is_location_input(text):
#     text = text.lower().strip()
#     if re.match(r"^\d{6}$", text):
#         return True
#     keywords = ["nagar", "layout", "road", "area", "sector", "block", "city", "town", "village", "locality"]
#     known_cities = {"bengaluru", "bangalore", "mumbai", "delhi", "chennai", "hyderabad", "pune", "noida", "gurgaon", "mysuru"}
#     return any(k in text for k in keywords) or any(word in known_cities for word in text.split())

# # === 💰 Detect budget input ===
# def is_budget_input(text):
#     text = text.lower().replace(",", "").strip()
#     if any(unit in text for unit in ["sqft", "sqm", "m2", "square", "meter"]):
#         return False
#     return bool(re.search(r"(₹)?\s?\d{2,7}\s?(k|l|cr|lakhs|crore)?", text))

# # === 📏 Detect square footage input ===
# def is_square_footage_input(text):
#     text = text.lower().replace(" ", "").strip()
#     return bool(re.match(r"^\d{2,5}(sqft|sq\.ft|sqm|m2|m²)?$", text))

# # === 🔍 Match partial option input ===
# def match_partial_response(prompt_response, user_input):
#     user_input = user_input.strip().lower()
#     for line in prompt_response.splitlines():
#         clean = line.lower().strip()
#         if user_input in clean:
#             return line
#         if re.match(rf"^{user_input}[\).\s]", clean):
#             return line
#     return None

# # === 💬 Main chatbot function ===
# def hybrid_chatbot(prompt_map, conversation, llm):
#     print("\n🤖 Chatbot is live. Type your input or 'exit' to quit.\n")
#     last_response = None

#     while True:
#         user_input = input("You: ").strip()
#         user_input_clean = user_input.lower()

#         if user_input_clean == "exit":
#             print("Bot: Goodbye!")
#             break

#         # === Handle state-based flow ===
#         if not user_state["Intent"]:
#             response = prompt_map.get("buyer")
#             print("\nBot:\n" + response)
#             user_state["Intent"] = user_input_clean
#             continue

#         if not user_state["User_type"]:
#             response = prompt_map.get("individual")
#             print("\nBot:\n" + response)
#             user_state["User_type"] = user_input_clean
#             continue

#         if not user_state["property_type"]:
#             response = prompt_map.get("just exploring for now")
#             print("\nBot:\n" + response)
#             user_state["property_type"] = user_input_clean
#             continue

#         if not user_state["sub_type"]:
#             response = prompt_map.get("both")
#             print("\nBot:\n" + response)
#             user_state["sub_type"] = user_input_clean
#             continue

#         if not user_state["location"] and is_location_input(user_input):
#             user_state["location"] = user_input
#             print("\nBot: Got it! Could you also share your budget range for this location?")
#             continue

#         if not user_state["budget"] and is_budget_input(user_input):
#             user_state["budget"] = user_input
#             print("\nBot: Great! How many bedrooms or rooms do you need?")
#             print("1. 1 BHK / 1 Room\n2. 2 BHK\n3. 3 BHK\n4. 4 BHK or more\n5. Open layout / Not applicable")
#             continue

#         if not user_state["bhk"] and any(k in user_input.lower() for k in ["bhk", "1", "2", "3", "4", "open"]):
#             user_state["bhk"] = user_input
#             print("\nBot: Do you have a minimum or preferred square footage in mind? (e.g., 1000 sqft, 150 sqm)")
#             continue

#         if not user_state["sqft"] and is_square_footage_input(user_input):
#             user_state["sqft"] = user_input
#             print("\nBot: Do you prefer the property to be:")
#             print("1. Fully Furnished\n2. Semi-Furnished\n3. Unfurnished\n4. No preference")
#             continue

#         if not user_state["furnishing"] and any(k in user_input.lower() for k in ["furnished", "semi", "unfurnished", "no preference", "1", "2", "3", "4"]):
#             user_state["furnishing"] = user_input
#             print("\nBot: Thanks! Here's a summary of what you've shared so far:")
#             for k, v in user_state.items():
#                 print(f"- {k.replace('_', ' ').title()}: {v}")
#             print("\nWould you like to proceed to view matching listings or speak with an agent?")
#             continue

#         # === Default fallback ===
#         response = prompt_map.get(user_input_clean)
#         if response:
#             print("\nBot:\n" + response)
#             last_response = response
#             continue

#         if last_response:
#             partial_match = match_partial_response(last_response, user_input)
#             if partial_match:
#                 print(f"\nBot: Noted: {partial_match}")
#                 continue

#         # === Final fallback to open chat ===
#         try:
#             llm_response = conversation.predict(input=user_input)
#             print("\nBot:", llm_response.strip())
#         except Exception as e:
#             print("❌ LLM Chat Error:", str(e))
#             print("Bot: Sorry, something went wrong.")

# # === 🚀 Run ===
# if __name__ == "__main__":
#     docx_path = "C:/Users/rithv/WORK/Projects/realty2/prompt_table.docx"
#     prompt_map = load_prompt_response_table(docx_path)
#     conversation = create_llm_chatbot()
#     llm = conversation.llm
#     hybrid_chatbot(prompt_map, conversation, llm)



import os
import re
from docx import Document
from langchain.memory import ConversationBufferMemory
from langchain.chains import ConversationChain
from langchain_community.chat_models import ChatOpenAI

os.environ["OPENAI_API_KEY"] = ""  # Replace with your actual key

user_state = {
    "intent": None,
    "user_type": None,
    "looking_to": None,
    "property_type": None,
    "sub_type": None,
    "location": None,
    "budget": None,
    "bhk": None,
    "sqft": None,
    "furnishing": None,
    "timeline": None,
    "current_question": "intent"
}

question_sequence = [
    "intent",
    "user_type",
    "looking_to",
    "property_type",
    "sub_type",
    "location",
    "budget",
    "bhk",
    "sqft",
    "furnishing",
    "timeline"
]

question_prompts = {
    "intent": "buyer",  # initial trigger
    "user_type": "buyer",  # triggers: Are you an Individual / Real Estate Agent
    "looking_to": "individual",  # triggers: Buy / Rent / Explore
    "property_type": "just exploring for now",
    "sub_type": "both",
    "bhk": '''1. 1 BHK / 1 Room
2. 2 BHK
3. 3 BHK
4. 4 BHK or more
5. Open layout / Not applicable''',
    "furnishing": '''1. Fully Furnished
2. Semi-Furnished
3. Unfurnished
4. No preference'''
}


def load_prompt_response_table(docx_path):
    doc = Document(docx_path)
    prompt_map = {}
    for table in doc.tables:
        for row in table.rows[1:]:
            prompt = row.cells[0].text.strip().lower()
            response = "\n".join(
                para.text.strip() for para in row.cells[1].paragraphs if para.text.strip()
            )
            prompt_map[prompt] = response
    return prompt_map


def create_llm_chatbot():
    llm = ChatOpenAI(model="gpt-3.5-turbo", temperature=0.5)
    memory = ConversationBufferMemory()
    return ConversationChain(llm=llm, memory=memory, verbose=False)


def is_location_input(text):
    text = text.lower().strip()
    if re.match(r"^\d{6}$", text):
        return True
    keywords = ["nagar", "layout", "road", "area", "sector", "block", "city", "town", "village", "locality"]
    known_cities = {"bengaluru", "bangalore", "mumbai", "delhi", "chennai", "hyderabad", "pune", "noida", "gurgaon", "mysuru"}
    return any(k in text for k in keywords) or any(word in known_cities for word in text.split())


def is_budget_input(text):
    text = text.lower().replace(",", "").strip()
    if any(unit in text for unit in ["sqft", "sqm", "m2", "square", "meter"]):
        return False
    return bool(re.search(r"(₹)?\s?\d{2,7}\s?(k|l|cr|lakhs|crore)?", text))


def is_square_footage_input(text):
    text = text.lower().replace(" ", "").strip()
    return bool(re.match(r"^\d{2,5}(sqft|sq\.ft|sqm|m2|m²)?$", text))


def match_partial_response(prompt_response, user_input):
    user_input = user_input.strip().lower()
    for line in prompt_response.splitlines():
        clean = line.lower().strip()
        if user_input in clean:
            return line
        if re.match(rf"^{user_input}[\).\s]", clean):
            return line
    return None


def get_next_question_key():
    for q in question_sequence:
        if user_state[q] is None:
            return q
    return None


def ask_question(q_key, prompt_map):
    print("Bot:")
    prompt_label = question_prompts.get(q_key)
    if prompt_label and prompt_label in prompt_map:
        print(prompt_map[prompt_label])
    elif q_key in ["bhk", "furnishing"]:
        print(question_prompts[q_key])
    elif q_key == "timeline":
        print("When are you planning to move or make a purchase? (e.g., Immediately, Within 1 month, 3–6 months, etc.)")
    else:
        print(f"Please tell me your {q_key.replace('_', ' ')}.")


def hybrid_chatbot(prompt_map, conversation, llm):
    print("\n🤖 Chatbot is live. Type your input or 'exit' to quit.\n")
    last_response = None

    user_state["intent"] = "buyer"
    user_state["current_question"] = get_next_question_key()
    ask_question(user_state["current_question"], prompt_map)

    while True:
        user_input = input("You: ").strip()
        user_input_clean = user_input.lower()
        q_key = user_state["current_question"]

        if user_input_clean == "exit":
            print("Bot: Goodbye!")
            break

        if q_key == "location" and is_location_input(user_input):
            user_state[q_key] = user_input

        elif q_key == "sub_type":
            property_type_raw = user_state.get("property_type", "").lower()
            property_type = re.sub(r"\d+[\).:\s]*", "", property_type_raw).strip()

            if property_type in prompt_map:
                print("Bot:")
                print(prompt_map[property_type])
                subtype_input = input("You: ").strip()
                matched = match_partial_response(prompt_map[property_type], subtype_input)
                user_state[q_key] = matched or subtype_input
            else:
                user_state[q_key] = user_input

        elif q_key in ["bhk", "furnishing"]:
            matched = match_partial_response(question_prompts[q_key], user_input)
            user_state[q_key] = matched or user_input
        else:
            prompt_label = question_prompts.get(q_key, "")
            matched_prompt = match_partial_response(prompt_map.get(prompt_label, ""), user_input)
            user_state[q_key] = matched_prompt or user_input

        next_q = get_next_question_key()
        user_state["current_question"] = next_q

        if next_q:
            ask_question(next_q, prompt_map)
        else:
            print("\nBot: Thanks! Here's a summary of what you've shared so far:")
            for k, v in user_state.items():
                if k != "current_question":
                    print(f"- {k.replace('_', ' ').title()}: {v}")
            print("\nWould you like to proceed to view matching listings or speak with an agent?")
            break


if __name__ == "__main__":
    docx_path = "prompt_table.docx"
    prompt_map = load_prompt_response_table(docx_path)
    conversation = create_llm_chatbot()
    llm = conversation.llm
    hybrid_chatbot(prompt_map, conversation, llm)
